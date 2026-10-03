"""Authoritative result access: paper-reported constants and sheet parsing.

Two distinct evidence sources live here:

1. ``PAPER_*`` constants — transcribed from the published paper PDF
   (``results/publication/BIOINFORMATICS_2026_398_CR.pdf``).  These are the
   authoritative published numbers.
2. ``parse_*`` functions — read the historical performance workbooks
   (``results/sheets/**``), which hold raw per-configuration metrics computed
   on *transformed* (normalised/log) targets.  They document the experiment
   process; they do not contain the paper's headline summary.

The historical workbook layout (see ``research_archive/src/find_best_models.py``
and ``research_archive/src/read_train.py``): each sheet is one pipeline
configuration with B1/B2/B3 holding train:test ratio, DR technique, and
normalisation method.  Rows 7-24 (skipping 12 and 21) hold the 16 models.
Columns A-H are the 90%-variance block and I-P the 95%-variance block, each
laid out as: model, search method, best params, MSE, MAE, MAPE, RMSE, SMAPE —
all computed on the transformed target of that configuration's test split.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

PAPER_PDF = "results/publication/BIOINFORMATICS_2026_398_CR.pdf"
PAPER_DOI = "10.5220/0014636000004070"
PAPER_VENUE = (
    "Proceedings of the 19th International Joint Conference on Biomedical "
    "Engineering Systems and Technologies (BIOSTEC) - Volume 2: BIOINFORMATICS"
)
PAPER_PAGES = "704-715"
# Table 2, PDF p. 11 (proceedings p. 714): "Error Metrics from the Un-normalized
# Test Data on an Hourly Scale".  Table 1, PDF p. 10 (proceedings p. 713) holds
# the same models' metrics computed on the normalised target instead.
PAPER_TABLE2_HOURLY = [
    {
        "method": "SOTA 1: non-temporal encoding", "region": "BA11", "pct_train": 80,
        "norm": "Minmax", "window": None, "dr": "PCA", "dr_level": 90, "model": "LSTM",
        "mae": 2.425, "mse": 9.482, "rmse": 3.079, "mape": 0.138, "smape": 13.361,
        "stderr": 3.077,
    },
    {
        "method": "SOTA 1: non-temporal encoding", "region": "BA47", "pct_train": 80,
        "norm": "Minmax", "window": None, "dr": "PCA", "dr_level": 90, "model": "LSTM",
        "mae": 3.274, "mse": 14.653, "rmse": 3.828, "mape": 0.194, "smape": 19.330,
        "stderr": 3.823,
    },
    {
        "method": "SOTA 2: temporal encoding via CNN", "region": "BA11", "pct_train": 70,
        "norm": "Minmax", "window": 3, "dr": "PCA", "dr_level": 90, "model": "Bagging Regressor",
        "mae": 0.945, "mse": 1.285, "rmse": 1.134, "mape": 0.069, "smape": 7.015,
        "stderr": 1.107,
    },
    {
        "method": "SOTA 2: temporal encoding via CNN", "region": "BA47", "pct_train": 80,
        "norm": "Log", "window": 3, "dr": "KPCA", "dr_level": 95, "model": "AdaBoost Regressor",
        "mae": 1.757, "mse": 5.945, "rmse": 2.438, "mape": 0.103, "smape": 10.887,
        "stderr": 2.201,
    },
    {
        "method": "Ours: temporal encoding via AutoEncoder", "region": "BA11",
        "pct_train": 80, "norm": "Minmax", "window": 3, "dr": "Isomap", "dr_level": 90,
        "model": "ExtraTrees Regressor",
        "mae": 0.839, "mse": 1.013, "rmse": 1.006, "mape": 0.055, "smape": 5.386,
        "stderr": 0.996,
    },
    {
        "method": "Ours: temporal encoding via AutoEncoder", "region": "BA47",
        "pct_train": 70, "norm": "Minmax", "window": 3, "dr": "PCA", "dr_level": 95,
        "model": "AdaBoost Regressor",
        "mae": 1.227, "mse": 2.153, "rmse": 1.467, "mape": 0.112, "smape": 9.778,
        "stderr": 1.451,
    },
]

MODEL_ROWS = [7, 8, 9, 10, 11, 13, 14, 15, 16, 17, 18, 19, 20, 22, 23, 24]  # 1-indexed
VARIANCE_BLOCKS = {90: 1, 95: 9}  # 1-indexed first column (A vs I) of each block
BLOCK_LAYOUT = ["model", "search", "best_params", "mse", "mae", "mape", "rmse", "smape"]

METRIC_COLUMNS = ["mse", "mae", "mape", "rmse", "smape"]


def paper_table2_frame() -> pd.DataFrame:
    """The paper's hourly-scale results (Table 2) as a DataFrame."""
    return pd.DataFrame(PAPER_TABLE2_HOURLY)


def check_paper_internal_consistency(rtol: float = 0.02) -> list[str]:
    """Sanity-check the transcribed constants; return a list of problems (empty = OK)."""
    problems: list[str] = []
    for row in PAPER_TABLE2_HOURLY:
        if not np.isclose(row["rmse"], np.sqrt(row["mse"]), rtol=rtol, atol=0.005):
            problems.append(
                f"{row['region']} {row['model']}: rmse {row['rmse']} != sqrt(mse {row['mse']})"
            )
    ours = [r for r in PAPER_TABLE2_HOURLY if r["method"].startswith("Ours")]
    for region in ("BA11", "BA47"):
        region_rows = [r for r in PAPER_TABLE2_HOURLY if r["region"] == region]
        best = min(region_rows, key=lambda r: r["mae"])
        if not best["method"].startswith("Ours"):
            problems.append(f"{region}: our method should have the lowest hourly MAE")
        ours_region = next(r for r in ours if r["region"] == region)
        if not np.isclose(ours_region["mae"], best["mae"]):
            problems.append(f"{region}: transcribed MAE disagrees with table minimum")
    return problems


def parse_performance_workbook(path: str | Path) -> list[dict]:
    """Parse one historical performance workbook into a list of record dicts."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    df_name_match = Path(path).name.split(" Overall Model Peformance Results")[0]
    records: list[dict] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        ratio = sheet["B1"].value
        dr_technique = sheet["B2"].value
        norm_method = sheet["B3"].value
        for row in MODEL_ROWS:
            for variance, first_col in VARIANCE_BLOCKS.items():
                values = [
                    sheet.cell(row=row, column=first_col + offset).value
                    for offset in range(len(BLOCK_LAYOUT))
                ]
                if all(v is None for v in values[3:]):
                    continue
                record = {
                    "workbook": df_name_match,
                    "sheet": sheet_name,
                    "train_test_ratio": ratio,
                    "dr_method": dr_technique,
                    "norm_method": norm_method,
                    "variance_level": variance,
                    "excel_row": row,
                }
                record.update(dict(zip(BLOCK_LAYOUT, values)))
                records.append(record)
    return records


def parse_sheets_tree(sheets_dir: str | Path) -> pd.DataFrame:
    """Parse every workbook under ``sheets_dir`` (including option/window folders)."""
    all_records: list[dict] = []
    for workbook_path in sorted(Path(sheets_dir).rglob("*Overall Model Peformance Results.xlsx")):
        relative = workbook_path.relative_to(sheets_dir).parts
        option = relative[0] if len(relative) > 1 else ""
        window = ""
        for part in relative:
            if part.startswith("window"):
                window = part
        for record in parse_performance_workbook(workbook_path):
            record["option"] = option
            record["window"] = window
            all_records.append(record)
    return pd.DataFrame.from_records(all_records)


def best_by_test_mae(tidy: pd.DataFrame) -> pd.DataFrame:
    """Best configuration per (option, window, workbook=region) by test MAE.

    .. warning::
       The parsed metrics come from the historical sheets and are computed on
       *transformed* targets whose scale differs between normalisation methods,
       so cross-method comparison is not meaningful (see docs/LIMITATIONS.md).
       This function reproduces the historical selection step for documentation.
    """
    if tidy.empty:
        return tidy
    return (
        tidy.sort_values("mae")
        .groupby(["option", "window", "workbook"], dropna=False)
        .head(1)
        .reset_index(drop=True)
    )
