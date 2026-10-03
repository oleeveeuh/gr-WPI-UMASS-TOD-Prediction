"""Parsing of the authoritative performance workbooks (synthetic replica)."""

import numpy as np
import openpyxl
import pytest

from tod_pred.authoritative import (
    MODEL_ROWS,
    check_paper_internal_consistency,
    parse_performance_workbook,
)


@pytest.fixture
def synthetic_workbook(tmp_path):
    """A minimal replica of the historical 'Overall Model Peformance Results' layout."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Option 1 - LR1 - DN1 (80-20)"
    sheet["A1"] = "Train:Test Ratio"
    sheet["B1"] = "80:20"
    sheet["A2"] = "Dimension Reduction Techniques"
    sheet["B2"] = "ICA"
    sheet["A3"] = "Data Normalization Method"
    sheet["B3"] = "MinMax"
    # First model row (Excel row 7), 90%-variance block (columns A-H).
    row_values = ["Linear Regression", "Grid", "{'fit_intercept': True}", 0.5, 0.4, 12.0, 0.7, 11.0]
    for offset, value in enumerate(row_values):
        sheet.cell(row=7, column=1 + offset, value=value)
    # 95%-variance block (columns I-P), same row.
    for offset, value in enumerate(["Linear Regression", "Grid", "{}", 0.9, 0.8, 20.0, 1.0, 19.0]):
        sheet.cell(row=7, column=9 + offset, value=value)
    path = tmp_path / "BA11 Overall Model Peformance Results.xlsx"
    workbook.save(path)
    return path


def test_parse_extracts_both_variance_blocks(synthetic_workbook):
    records = parse_performance_workbook(synthetic_workbook)
    assert len(records) == 2  # one model row x two variance blocks
    by_variance = {record["variance_level"]: record for record in records}
    assert by_variance[90]["mae"] == pytest.approx(0.4)
    assert by_variance[95]["mae"] == pytest.approx(0.8)
    assert by_variance[90]["dr_method"] == "ICA"
    assert by_variance[90]["norm_method"] == "MinMax"
    assert by_variance[90]["model"] == "Linear Regression"


def test_parse_skips_empty_rows(synthetic_workbook):
    records = parse_performance_workbook(synthetic_workbook)
    assert all(record["excel_row"] in MODEL_ROWS for record in records)
    # Only row 7 had values; the other 15 model rows yielded no records.


def test_paper_constants_are_internally_consistent():
    assert check_paper_internal_consistency() == []


def test_paper_headline_values():
    from tod_pred.authoritative import paper_table2_frame

    table = paper_table2_frame()
    ours = table[table["method"].str.startswith("Ours")]
    assert set(ours["model"]) == {"ExtraTrees Regressor", "AdaBoost Regressor"}
    np.testing.assert_allclose(sorted(ours["mae"]), [0.839, 1.227])
